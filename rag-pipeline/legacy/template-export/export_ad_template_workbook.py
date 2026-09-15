"""Legacy: export NH advertisement-template HWPX into an auditable review workbook.

The HWPX is a template aid, not a regulation source.  The exporter preserves
source text and shows only exact v2 section/label matches as candidates.
Decision-guide fields are deliberately empty: they must not be invented.
"""
from __future__ import annotations

import argparse
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


K = {
    "source_table": "\uc6d0\ubcf8\ud45c\ubc88\ud638", "source_row": "\uc6d0\ubcf8\ud589\ubc88\ud638",
    "template_id": "\ud15c\ud50c\ub9bf_\uc6d0\ubcf8ID", "section": "\ud15c\ud50c\ub9bf\uc139\uc158",
    "label": "\uad6c\ubd84", "example": "\uc608\uc2dc\ubb38\uad6c", "required": "\ud544\uc218\uc5ec\ubd80",
    "guidance": "\uae30\uc7ac\uc694\ub839", "image": "\uc774\ubbf8\uc9c0_\ub300\uccb4\ubb38\uad6c\ud3ec\ud568",
    "candidate": "v2_T\ud56d\ubaa9ID_\ud6c4\ubcf4", "mapping": "v2_\ub9e4\ud551\uc0c1\ud0dc",
    "ok": "\ud310\uc815\uac00\uc774\ub4dc_\uc801\ud569\uc870\uac74", "bad": "\ud310\uc815\uac00\uc774\ub4dc_\ubd80\uc801\uc815\uc870\uac74",
    "review": "\ud310\uc815\uac00\uc774\ub4dc_\ud655\uc778\ud544\uc694\uc870\uac74", "notice": "\ud310\uc815\uac00\uc774\ub4dc_\uc548\ub0b4\ubb38\uad6c",
}
RAW, SUMMARY, REVIEW = "\ud15c\ud50c\ub9bf_\uc6d0\ubb38\uc815\ub9ac", "\ud15c\ud50c\ub9bf_\uc694\uc57d", "\ucd94\ucd9c_\uac80\ud1a0\ud544\uc694"
YES, TRIANGLE = "O", "\u25b3"


def lname(node): return node.tag.rsplit("}", 1)[-1]
def clean(value): return re.sub(r"[ \t]+", " ", "" if value is None else str(value).replace("\xa0", " ")).strip()
def norm(value): return re.sub(r"\s+", "", clean(value)).replace("[", "").replace("]", "")


def cell_text(cell):
    values = []
    for p in cell.iter():
        if lname(p) == "p":
            value = clean("".join(p.itertext()))
            if value: values.append(value)
    return "\n".join(dict.fromkeys(values))


def read_table(table):
    cells = []
    for cell in table.iter():
        if lname(cell) != "tc": continue
        address = next((n for n in cell if lname(n) == "cellAddr"), None)
        span = next((n for n in cell if lname(n) == "cellSpan"), None)
        if address is None: continue
        rs = int(span.get("rowSpan", "1")) if span is not None else 1
        cs = int(span.get("colSpan", "1")) if span is not None else 1
        cells.append((int(address.get("rowAddr", "0")), int(address.get("colAddr", "0")), rs, cs, cell_text(cell)))
    row_count, col_count = int(table.get("rowCnt", "0")), int(table.get("colCnt", "0"))
    def at(row, col):
        for r, c, rs, cs, value in cells:
            if r <= row < r + rs and c <= col < c + cs: return value
        return ""
    matrix = [[at(row, col) for col in range(col_count)] for row in range(row_count)]
    headers = [clean(value) for value in matrix[0]] if matrix else []
    expected = {norm(value) for value in (K["label"], K["example"], K["required"], K["guidance"])}
    missing = expected - {norm(value) for value in headers}
    return headers, matrix[1:], (["\ud45c \ud5e4\ub354 \ub204\ub77d \ub610\ub294 \ubcc0\ud615: " + ", ".join(sorted(missing))] if missing else [])


def find_col(headers, name): return next((i for i, x in enumerate(headers) if norm(x) == norm(name)), None)


def extract(path):
    with zipfile.ZipFile(path) as archive: root = ET.fromstring(archive.read("Contents/section0.xml"))
    rows, issues, section, table_no = [], [], "", 0
    for paragraph in list(root):
        if lname(paragraph) != "p": continue
        tables = [node for node in paragraph.iter() if lname(node) == "tbl"]
        if not tables:
            heading = clean("".join(n.text or "" for n in paragraph.iter() if lname(n) == "t"))
            if re.fullmatch(r"\[[^\]]+\]", heading): section = heading[1:-1].strip()
            continue
        for table in tables:
            table_no += 1; headers, source_rows, table_issues = read_table(table)
            issues.extend({K["source_table"]: str(table_no), K["section"]: section, "\uac80\ud1a0\uc0ac\uc720": x} for x in table_issues)
            idx = {field: find_col(headers, K[field]) for field in ("label", "example", "required", "guidance")}
            for row_no, values in enumerate(source_rows, 1):
                def value(field):
                    index = idx[field]
                    return clean(values[index]) if index is not None and index < len(values) else ""
                row = {K["source_table"]:table_no, K["source_row"]:row_no, K["template_id"]:f"HWPX-T{table_no:02d}-R{row_no:03d}", K["section"]:section,
                       K["label"]:value("label"), K["example"]:value("example"), K["required"]:value("required"), K["guidance"]:value("guidance")}
                if not any(row[x] for x in (K["label"],K["example"],K["required"],K["guidance"])): continue
                row[K["image"]] = "Y" if "\uadf8\ub9bc\uc785\ub2c8\ub2e4" in row[K["example"]] else "N"; rows.append(row)
    return rows, issues


def candidates(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True); sheet = book["\uc2e0\uaddc\uaddc\uce59\ud6c4\ubcf4"]
    data = sheet.iter_rows(values_only=True); headers = [clean(x) for x in next(data)]; index = {x:i for i,x in enumerate(headers)}; output = defaultdict(list)
    for row in data:
        def get(name): return clean(row[index[name]]) if name in index and index[name] < len(row) else ""
        item, section, label = get("\ud15c\ud50c\ub9bf_\uc140ID"), get("\uc139\uc158"), get(K["label"])
        if item and section and label: output[(norm(section),norm(label))].append(item)
    book.close(); return output


def style(sheet, widths):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font, cell.fill = Font(color="FFFFFF",bold=True), fill; cell.alignment = Alignment(horizontal="center",vertical="center",wrap_text=True)
    sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
    for letter, width in widths.items(): sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical="top",wrap_text=True)


def build(output, source_path, rows, issues):
    book = Workbook(); readme = book.active; readme.title = "README"
    for row in [
        ("\ud56d\ubaa9","\ub0b4\uc6a9"),("\uc6d0\ubcf8",str(source_path)),("\uc0dd\uc131\uc77c",date.today().isoformat()),
        ("\ucd94\ucd9c \ubc94\uc704","HWPX \ud15c\ud50c\ub9bf \ud45c\uc758 \uad6c\ubd84, \uc608\uc2dc\ubb38\uad6c, \ud544\uc218\uc5ec\ubd80, \uae30\uc7ac\uc694\ub839\uc744 \ud589 \ub2e8\uc704\ub85c \ubcf4\uc874"),
        ("\ud310\uc815 \uac00\uc774\ub4dc","\uc801\ud569, \ubd80\uc801\uc815, \ud655\uc778\ud544\uc694, \uc548\ub0b4\ubb38\uad6c\ub294 \uc6d0\ubb38\uc5d0 \uc5c6\uc5b4 \ube44\uc6cc \ub454 \uac80\uc218 \uce78. \uac1c\ubcc4 \uad11\uace0\ub85c \ucd94\uc815\ud558\uc9c0 \uc54a\uc74c."),
        ("v2 \ub9e4\ud551","\uaddc\uc81c\ubaa9\ub85d v2 \uc2e0\uaddc\uaddc\uce59\ud6c4\ubcf4\uc758 \uc139\uc158+\uad6c\ubd84\uc774 \uc644\uc804\ud788 \uac19\uc740 \uacbd\uc6b0\ub9cc \ud6c4\ubcf4\ub85c \ud45c\uc2dc. \ubbf8\ub9e4\ud551\uc740 \uc624\ub958\uac00 \uc544\ub2cc \ucd94\uac00 \uac80\ud1a0 \ub300\uc0c1.")]: readme.append(row)
    style(readme,{"A":22,"B":120})
    raw = book.create_sheet(RAW); columns=[K[x] for x in ("source_table","source_row","template_id","section","label","example","required","guidance","image","candidate","mapping","ok","bad","review","notice")]; raw.append(columns)
    for row in rows: raw.append([row.get(x,"") for x in columns])
    style(raw,{"A":11,"B":11,"C":20,"D":36,"E":20,"F":70,"G":12,"H":55,"I":16,"J":20,"K":14,"L":45,"M":45,"N":45,"O":45})
    summary=book.create_sheet(SUMMARY); summary.append([K["section"],"\ud56d\ubaa9\uc218",YES,TRIANGLE,"\ud544\uc218\uc5ec\ubd80 \ubbf8\uae30\uc7ac","v2 \uc815\ud655\uc77c\uce58","v2 \ubcf5\uc218\ud6c4\ubcf4","v2 \ubbf8\ub9e4\ud551"]); groups=defaultdict(list)
    for row in rows: groups[row[K["section"]]].append(row)
    for section, items in sorted(groups.items()):
        statuses, reqs = Counter(x[K["mapping"]] for x in items), Counter(x[K["required"]] for x in items)
        summary.append([section,len(items),reqs[YES],reqs[TRIANGLE],len(items)-reqs[YES]-reqs[TRIANGLE],statuses["\uc815\ud655\uc77c\uce58"],statuses["\ubcf5\uc218\ud6c4\ubcf4"],statuses["\ubbf8\ub9e4\ud551"]])
    style(summary,{"A":48,"B":10,"C":9,"D":9,"E":15,"F":13,"G":13,"H":12})
    review=book.create_sheet(REVIEW); review.append([K["source_table"],K["section"],"\uac80\ud1a0\uc0ac\uc720"])
    for issue in issues: review.append([issue[K["source_table"]],issue[K["section"]],issue["\uac80\ud1a0\uc0ac\uc720"]])
    for row in rows:
        if row[K["mapping"]] == "\ubcf5\uc218\ud6c4\ubcf4": review.append([row[K["source_table"]],row[K["section"]],f"{row[K['template_id']]}: v2 \ubcf5\uc218\ud6c4\ubcf4 ({row[K['candidate']]})"])
    style(review,{"A":13,"B":48,"C":90}); output.parent.mkdir(parents=True,exist_ok=True); book.save(output)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--template-hwpx",type=Path,required=True); parser.add_argument("--regulation-v2",type=Path,required=True); parser.add_argument("--output",type=Path,required=True); args=parser.parse_args()
    rows, issues=extract(args.template_hwpx); pool=candidates(args.regulation_v2)
    for row in rows:
        found=pool.get((norm(row[K["section"]]),norm(row[K["label"]])),[]); row[K["candidate"]]="\n".join(found); row[K["mapping"]]="\uc815\ud655\uc77c\uce58" if len(found)==1 else ("\ubcf5\uc218\ud6c4\ubcf4" if found else "\ubbf8\ub9e4\ud551")
        for x in ("ok","bad","review","notice"): row[K[x]]=""
    build(args.output,args.template_hwpx,rows,issues); print(f"template rows={len(rows)} tables={len(set(x[K['source_table']] for x in rows))} output={args.output}")


if __name__ == "__main__": main()
