"""Legacy: export every textual element of the NH template HWPX to Excel.

No regulation mapping or inferred judgement fields are added.  The workbook
contains a friendly row view, the physical HWPX cells (including merge spans),
and all text paragraphs outside tables so extraction can be audited.
"""
from __future__ import annotations

import argparse
import re
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def local_name(node: ET.Element) -> str:
    return node.tag.rsplit("}", 1)[-1]


def clean(value: object) -> str:
    text = "" if value is None else str(value).replace("\u00a0", " ")
    return re.sub(r"[ \t]+", " ", text).strip()


def paragraph_text(node: ET.Element) -> str:
    return clean("".join(item.text or "" for item in node.iter() if local_name(item) == "t"))


def cell_text(cell: ET.Element) -> str:
    values: list[str] = []
    for paragraph in cell.iter():
        if local_name(paragraph) != "p":
            continue
        value = paragraph_text(paragraph)
        if value:
            values.append(value)
    return "\n".join(dict.fromkeys(values))


def physical_cells(table: ET.Element) -> list[dict[str, object]]:
    result = []
    for cell in table.iter():
        if local_name(cell) != "tc":
            continue
        address = next((item for item in cell if local_name(item) == "cellAddr"), None)
        span = next((item for item in cell if local_name(item) == "cellSpan"), None)
        if address is None:
            continue
        result.append({
            "row": int(address.get("rowAddr", "0")),
            "col": int(address.get("colAddr", "0")),
            "row_span": int(span.get("rowSpan", "1")) if span is not None else 1,
            "col_span": int(span.get("colSpan", "1")) if span is not None else 1,
            "text": cell_text(cell),
        })
    return result


def expanded_matrix(table: ET.Element, cells: list[dict[str, object]]) -> list[list[str]]:
    row_count = int(table.get("rowCnt", "0"))
    col_count = int(table.get("colCnt", "0"))
    matrix = [["" for _ in range(col_count)] for _ in range(row_count)]
    for cell in cells:
        for row in range(int(cell["row"]), int(cell["row"]) + int(cell["row_span"])):
            for col in range(int(cell["col"]), int(cell["col"]) + int(cell["col_span"])):
                if row < row_count and col < col_count:
                    matrix[row][col] = str(cell["text"])
    return matrix


def extract(path: Path):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("Contents/section0.xml"))

    tables, cells, outside = [], [], []
    section = ""
    paragraph_number = 0
    table_number = 0
    for paragraph in list(root):
        if local_name(paragraph) != "p":
            continue
        paragraph_number += 1
        nested_tables = [item for item in paragraph.iter() if local_name(item) == "tbl"]
        if not nested_tables:
            text = paragraph_text(paragraph)
            if text:
                outside.append({"paragraph": paragraph_number, "section": section, "text": text})
                if re.fullmatch(r"\[[^\]]+\]", text):
                    section = text[1:-1].strip()
            continue
        for table in nested_tables:
            table_number += 1
            table_cells = physical_cells(table)
            matrix = expanded_matrix(table, table_cells)
            headers = matrix[0] if matrix else []
            tables.append({
                "table": table_number,
                "section": section,
                "rows": len(matrix),
                "cols": len(headers),
                "headers": headers,
                "matrix": matrix,
            })
            for cell in table_cells:
                cells.append({"table": table_number, "section": section, **cell})
    return tables, cells, outside


def style(sheet, widths: dict[str, int]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build(output: Path, source: Path, tables, cells, outside) -> None:
    book = Workbook()
    readme = book.active
    readme.title = "README"
    for row in [
        ["항목", "내용"],
        ["원본", str(source)],
        ["작성 원칙", "원본 HWPX의 표와 표 밖 문구만 수록. 규제목록 매핑·판정 조건·추정값을 추가하지 않음"],
        ["원본 빈칸", "원본 셀이 비어 있을 때만 빈칸으로 유지"],
        ["원본_행정리", "병합 셀 값을 하위 행에도 펼친 사람이 읽기 쉬운 보기"],
        ["원본_셀그대로", "실제 HWPX 셀의 행·열 주소와 병합 범위를 보존한 감사용 보기"],
        ["표밖_원문", "템플릿 제목과 표 밖 설명 문구"],
    ]:
        readme.append(row)
    style(readme, {"A": 22, "B": 110})

    summary = book.create_sheet("원본_표목록")
    summary.append(["원본표번호", "템플릿섹션", "원본행수(헤더포함)", "원본열수", "원본헤더"])
    for table in tables:
        summary.append([table["table"], table["section"], table["rows"], table["cols"], " | ".join(table["headers"])])
    style(summary, {"A": 12, "B": 48, "C": 18, "D": 12, "E": 80})

    max_columns = max((table["cols"] for table in tables), default=0)
    rows_sheet = book.create_sheet("원본_행정리")
    rows_sheet.append(["원본표번호", "템플릿섹션", "원본행번호"] + [f"원본열{i + 1}" for i in range(max_columns)])
    for table in tables:
        for row_number, values in enumerate(table["matrix"], start=1):
            rows_sheet.append([table["table"], table["section"], row_number] + values + [""] * (max_columns - len(values)))
    widths = {"A": 12, "B": 48, "C": 12}
    for index in range(max_columns):
        widths[chr(ord("D") + index)] = 65 if index == 1 else 28
    style(rows_sheet, widths)

    cells_sheet = book.create_sheet("원본_셀그대로")
    cells_sheet.append(["원본표번호", "템플릿섹션", "행주소(0부터)", "열주소(0부터)", "행병합수", "열병합수", "원본문구"])
    for cell in cells:
        cells_sheet.append([cell["table"], cell["section"], cell["row"], cell["col"], cell["row_span"], cell["col_span"], cell["text"]])
    style(cells_sheet, {"A": 12, "B": 48, "C": 14, "D": 14, "E": 12, "F": 12, "G": 95})

    outside_sheet = book.create_sheet("표밖_원문")
    outside_sheet.append(["원본문단번호", "직전템플릿섹션", "원본문구"])
    for item in outside:
        outside_sheet.append([item["paragraph"], item["section"], item["text"]])
    style(outside_sheet, {"A": 15, "B": 48, "C": 110})

    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    tables, cells, outside = extract(args.source)
    build(args.output, args.source, tables, cells, outside)
    print(f"tables={len(tables)} cells={len(cells)} outside_paragraphs={len(outside)} output={args.output}")


if __name__ == "__main__":
    main()
