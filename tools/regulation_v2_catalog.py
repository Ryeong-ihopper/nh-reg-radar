# -*- coding: utf-8 -*-
"""Load regulation-v2 candidate items without dataset or answer-key state."""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from rag import build_items as v2_source  # noqa: E402


def product_groups_from_template_section(section: str) -> list[str]:
    """템플릿 섹션의 상품군을 넓히지 않고 그대로 보존한다."""
    for product in ("대출성", "예금성", "카드", "투자성"):
        if section.startswith(product):
            return [product]
    if section.startswith("전체"):
        return ["전체"]
    return []


def load_template_candidate_rules() -> list[dict[str, Any]]:
    book = openpyxl.load_workbook(v2_source.AGENT, read_only=True, data_only=True)
    sheet = book["신규규칙후보"]
    header = [cell.value for cell in next(sheet.iter_rows())]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, values))
        item_id = str(record.get("템플릿_셀ID") or "").strip()
        if not item_id:
            continue
        mandatory = str(record.get("필수여부") or "").strip()
        example = str(record.get("예시문구") or "").strip()
        guide = str(record.get("기재요령") or "").strip()
        section = str(record.get("섹션") or "").strip()
        product_groups = product_groups_from_template_section(section)
        rows.append({
            "item_id": item_id,
            "category": "PRESENCE",
            "category_label": "표시의무(신규규칙후보)",
            "title": f"{record.get('구분') or '템플릿 항목'} 표시",
            "question": f"v2 T 규칙의 템플릿 항목 '{record.get('구분')}'이 이 광고에 적용되며 의미상 충족되는가?",
            "criterion": (
                f"사용자가 사용 가능하다고 확정한 v2 T 규칙이다. "
                f"필수여부={mandatory or '미기재'}. 예시={example} {guide} "
                f"O는 적용 템플릿에서 필수로 검사한다. △ 또는 조건부 항목은 "
                f"기재요령의 조건이 확인될 때만 필수이며, 조건을 확인할 수 없으면 UNDETERMINED다."
            ),
            "product_groups": product_groups,
            "product_subtype": section,
            "judgment_mode": "템플릿규칙",
            "judgment_type": "O필수·△조건부",
            "input_requirement": "광고물+템플릿 적용성",
            "required_medium": "v2 기재값",
            "violation_grade": None,
            "legal_basis": None,
            "representative_rule_id": None,
            "supporting_rules": [],
            "rule_summaries": [],
            "basis_details": [],
            "v2_note": "사용자가 판정 규칙으로 사용 가능하다고 확정",
            "source_sheet": "신규규칙후보",
            "template_required": mandatory,
            "example_text": example,
            "guide": guide,
        })
    return rows
