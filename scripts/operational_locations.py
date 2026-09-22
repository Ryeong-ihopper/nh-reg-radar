"""Saved evidence -> display geometry. Never use fuzzy text similarity."""
from __future__ import annotations

import math
import copy
import hashlib
import re
import sys
from functools import lru_cache
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "rag-pipeline"))
from rag.judgment.reading_quality import (VERSION as READING_QUALITY_VERSION, apply_reading_guard,
                                          needs_reading_review, project_reading_citations)  # noqa: E402
from rag.judgment.source_checks import (observed_grounding_errors, unresolved_applicability,
                                        template_heading_only_citation)  # noqa: E402
from rag.judgment.arithmetic import calculate_loan_rates, METHOD as ARITHMETIC_METHOD  # noqa: E402
from rag.judgment.grounding import cited_window_text, ungrounded_source_quotes  # noqa: E402
from rag.templates.catalog import required_observation_medium  # noqa: E402


def load_template_appropriate_judgments(path):
    """Load display-only guidance by exact template section and row label.

    The guide never enters retrieval or model prompts. Ambiguous duplicate
    labels fail closed instead of attaching another row's guidance.
    """
    import openpyxl

    source = Path(path)
    section = re.sub(r"^\s*\d+\s*[._-]?\s*", "", source.stem).strip()
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    values = {}
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = None
            for row in rows:
                normalized = [str(value).strip() if value is not None else "" for value in row]
                if "구분" in normalized and "적정 판단" in normalized:
                    header = normalized
                    break
            if header is None:
                continue
            label_column = header.index("구분")
            judgment_column = header.index("적정 판단")
            for row in rows:
                label = str(row[label_column] or "").strip() if label_column < len(row) else ""
                judgment = str(row[judgment_column] or "").strip() if judgment_column < len(row) else ""
                if not label or not judgment:
                    continue
                key = (section, label)
                if key in values and values[key] != judgment:
                    raise ValueError(f"ambiguous appropriate judgment for {section} / {label}")
                values[key] = judgment
    finally:
        book.close()
    if not values:
        raise ValueError("template guide contains no '구분' / '적정 판단' rows")
    return values


def page_lines(page):
    for region in page.get("regions", []):
        yield from region.get("lines", [])
    yield from page.get("unassigned_lines", [])


def valid_box(box, width, height):
    values = [width, height, *(box if isinstance(box, (list, tuple)) else [])]
    return (
        len(values) == 6
        and all(isinstance(v, (int, float)) and not isinstance(v, bool)
                and math.isfinite(v) for v in values)
        and width > 0 and height > 0
        and 0 <= box[0] < box[2] <= width
        and 0 <= box[1] < box[3] <= height
    )


def page_asset(document, page):
    """Global combined page and original asset-local page are different IDs."""
    number = page["page_no"]
    assets = (document.get("diagnostics") or {}).get("asset_pages") or {}
    matches = [(key, span) for key, span in assets.items()
               if span["start"] <= number <= span["end"]]
    if len(matches) == 1:
        key, span = matches[0]
        return key, number - span["start"] + 1
    return None, number


def with_rendered_line_locations(document, layout, asset_id):
    """Display-only exact alignment for a single HWP asset's saved render.

    Keep semantic refs/pages immutable. Only whitespace may differ; a wrapped
    sentence must match a unique contiguous run within one visual region.
    Repeated semantic text or repeated visual matches remain unlinked.
    """
    if layout.get("source") != "HWP 200-DPI render + parser visual projection":
        return document
    assets = (document.get("diagnostics") or {}).get("asset_pages") or {}
    if len(assets) != 1 or asset_id not in assets:
        return document
    def normalize(text):
        return re.sub(r"\s+", "", str(text or ""))
    semantic = [(page, line) for page in document.get("pages", [])
                for line in page_lines(page) if line.get("line_ref")]
    counts = {}
    for _, line in semantic:
        text = normalize(line.get("text") or line.get("parser_text"))
        counts[text] = counts.get(text, 0) + 1
    mappings = {}
    for page, line in semantic:
        text = normalize(line.get("text") or line.get("parser_text"))
        if not text or counts[text] != 1 or valid_box(
                line.get("bbox"), page.get("canvas_w"), page.get("canvas_h")):
            continue
        matches = []
        for visual_page in layout.get("pages", []):
            if visual_page.get("asset_id") not in (None, asset_id):
                continue
            for region in visual_page.get("regions", []):
                lines = region.get("lines", [])
                for start in range(len(lines)):
                    joined, run = "", []
                    for visual_line in lines[start:]:
                        part = normalize(visual_line.get("text"))
                        if not part or not valid_box(visual_line.get("bbox"),
                                visual_page.get("canvas_w"), visual_page.get("canvas_h")):
                            break
                        joined += part
                        run.append(visual_line)
                        if not text.startswith(joined):
                            break
                        if joined == text:
                            matches.append((visual_page, run))
                            break
        if len(matches) != 1:
            continue
        visual_page, run = matches[0]
        ref = line["line_ref"]
        mappings[ref] = [{
            "key": f"{ref}:render:{index}", "pageNo": visual_page["page_no"],
            "bbox": list(item["bbox"]), "width": visual_page["canvas_w"],
            "height": visual_page["canvas_h"], "asset_id": asset_id,
            "source_page_no": visual_page.get("source_page_no") or visual_page["page_no"],
            "precision": "LINE", "mapping_method": "EXACT_RENDERED_TEXT",
        } for index, item in enumerate(run)]
    return {**document, "display_line_locations": mappings}


def resolve_locations(document, judgment, evidence):
    """Exact parser refs or explicitly approximate source regions, never guessed lines."""
    judgment = project_reading_citations(judgment)
    pages = document.get("pages", [])
    index = {line["line_ref"]: (page, line) for page in pages
             for line in page_lines(page) if line.get("line_ref")}
    cited = [evidence[eid] for eid in judgment.get("evidence_ids", []) if eid in evidence]
    refs = list(dict.fromkeys(judgment.get("evidence_line_refs") or []))
    if not refs:
        # An evidence chunk containing several lines does not identify which
        # line supports the judgment.  Showing the whole chunk as one apparent
        # source range made unrelated text look like the cited sentence.  A
        # line fallback is safe only when the cited document has one exact line.
        refs = list(dict.fromkeys(ref for doc in cited
                    if doc.get("span_status") in {"parser_line_exact", "selected_text_line_aligned"}
                    and len(doc.get("line_refs") or []) == 1
                    for ref in doc.get("line_refs", [])))
    locations = []

    def append(page, box, key, precision):
        width, height = page.get("canvas_w"), page.get("canvas_h")
        if not valid_box(box, width, height):
            return
        asset, source_page = page_asset(document, page)
        locations.append({"key": key, "pageNo": page["page_no"], "bbox": box,
                          "width": width, "height": height, "asset_id": asset,
                          "source_page_no": source_page, "precision": precision})

    for ref in refs:
        if ref in index:
            page, line = index[ref]
            append(page, line.get("bbox"), ref, "LINE")
            if not valid_box(line.get("bbox"), page.get("canvas_w"), page.get("canvas_h")):
                locations.extend(copy.deepcopy(document.get("display_line_locations", {}).get(ref, [])))
    # A selected VLM passage with no proven line alignment can identify its
    # source region, but must not be presented as an exact parser-line quote.
    seen_regions = set()
    for doc in cited:
        if doc.get("span_status") != "region_level_selected_text":
            continue
        key = (doc.get("page_no"), doc.get("region_id"))
        if key in seen_regions:
            continue
        seen_regions.add(key)
        for page in pages:
            if page["page_no"] != key[0]:
                continue
            for region in page.get("regions", []):
                if region.get("region_id") == key[1]:
                    append(page, region.get("bbox"), f"region:{key[0]}:{key[1]}", "REGION")
    return locations


def resolve_chunk_locations(document, judgment, evidence):
    """Project each cited retrieval chunk as labeled context geometry.

    Exact judgment citations remain in ``evidence_locations``.  This separate
    projection lets the reviewer see the full text window supplied to the
    model without presenting every line in that window as the decisive quote.
    """
    cited_ids = list(dict.fromkeys(str(value) for value in judgment.get("evidence_ids") or []))
    chunks = []
    for evidence_id in cited_ids:
        doc = evidence.get(evidence_id)
        if not isinstance(doc, dict):
            continue
        refs = list(dict.fromkeys(str(value) for value in doc.get("line_refs") or []))
        locations = resolve_locations(
            document,
            {"evidence_ids": [evidence_id], "evidence_line_refs": refs},
            evidence,
        )
        grouped = {}
        for location in locations:
            group = (
                location["pageNo"], location.get("asset_id"),
                location.get("source_page_no"), location["width"], location["height"],
            )
            grouped.setdefault(group, []).append(location["bbox"])
        for number, (group, boxes) in enumerate(grouped.items()):
            page_no, asset_id, source_page_no, width, height = group
            chunks.append({
                "key": f"chunk:{evidence_id}:{page_no}:{number}",
                "pageNo": page_no,
                "bbox": [
                    min(box[0] for box in boxes), min(box[1] for box in boxes),
                    max(box[2] for box in boxes), max(box[3] for box in boxes),
                ],
                "width": width,
                "height": height,
                "asset_id": asset_id,
                "source_page_no": source_page_no,
                "precision": "CHUNK",
                "evidence_id": evidence_id,
            })
    return chunks


def resolve_review_locations(document, judgment, evidence):
    """Map uncertain source regions as review targets, never as judgment proof."""
    issues = (judgment.get("reading_quality_review") or {}).get("issues") or []
    ids = list(dict.fromkeys(str(value) for issue in issues
                            for value in issue.get("evidence_ids") or []))
    refs = list(dict.fromkeys(str(value) for issue in issues
                             for value in issue.get("line_refs") or []))
    if not ids and not refs:
        return []
    return resolve_locations(document, {"evidence_ids": ids, "evidence_line_refs": refs}, evidence)


def local_reading_review(document):
    """Return exact local parser-review targets without changing ad coverage."""
    locations, issues, seen = [], [], set()

    def append(page, value, key, precision, codes):
        width, height, box = page.get("canvas_w"), page.get("canvas_h"), value.get("bbox")
        if not valid_box(box, width, height) or key in seen:
            return
        seen.add(key)
        asset, source_page = page_asset(document, page)
        locations.append({"key": key, "pageNo": page["page_no"], "bbox": box,
                          "width": width, "height": height, "asset_id": asset,
                          "source_page_no": source_page, "precision": precision})
        issues.extend({"location": key, "code": code} for code in codes)

    for page in document.get("pages") or []:
        for number, region in enumerate(page.get("regions") or []):
            uncertain = needs_reading_review(region)
            # Missing final_text in legacy data is not evidence of an empty
            # region.  Current parser output explicitly writes the field.
            empty = "final_text" in region and not str(region.get("final_text") or "").strip()
            region_id = region.get("region_id") or number
            if empty:
                codes = ["EMPTY_LOCAL_REGION"]
                if uncertain:
                    codes.append("UNCERTAIN_LOCAL_READING")
                append(page, region, f"reading-region:{page['page_no']}:{region_id}", "REGION", codes)
            elif uncertain:
                before = len(locations)
                for line_number, line in enumerate(region.get("lines") or []):
                    line_ref = line.get("line_ref") or line_number
                    append(page, line, f"reading-line:{page['page_no']}:{line_ref}", "LINE",
                           ["UNCERTAIN_LOCAL_READING"])
                if len(locations) == before:
                    append(page, region, f"reading-region:{page['page_no']}:{region_id}", "REGION",
                           ["UNCERTAIN_LOCAL_READING"])
        for number, line in enumerate(page.get("unassigned_lines") or []):
            if needs_reading_review(line):
                line_ref = line.get("line_ref") or number
                append(page, line, f"reading-line:{page['page_no']}:{line_ref}", "LINE",
                       ["UNCERTAIN_LOCAL_READING"])
    return locations, issues


def saved_workspace(raw, requests, document, advertisement_id, discovery=None, reading_audits=None, rule_metadata=None):
    """Single projection shared by the active UI and JSON download."""
    ads = [ad for ad in raw.get("ads", []) if ad["ad_id"] == advertisement_id]
    scoped = {}
    reading_payloads = {}
    arithmetic = {}
    for request in requests:
        if request.get("ad_id") != advertisement_id and not request.get("ad_id", "").startswith(advertisement_id + "::"):
            continue
        key = request["ad_id"]
        definitions, evidence = scoped.setdefault(key, ({}, {}))
        definitions.update({rule["item_id"]: rule for rule in request["rules"]})
        evidence.update({doc["evidence_id"]: doc for doc in request["documents"]})
        for rule in request['rules']:
            reading_payloads[(key, rule['item_id'])] = request
            verified = calculate_loan_rates(request, rule)
            if verified:
                arithmetic[(key, rule['item_id'])] = verified
    texts = {line["line_ref"]: line.get("text") or line.get("parser_text") or ""
             for page in document.get("pages", []) for line in page_lines(page)
             if line.get("line_ref")}
    rows, review_rows = [], []
    for ad in ads:
        key = ad.get("scope_id") or ad["ad_id"]
        rules, evidence = scoped.get(key, ({}, {}))
        for candidates, target, kind in ((ad.get("candidates", []), rows, "finding"),
                                         (ad.get("review_candidates", []), review_rows, "review"),
                                         (ad.get("excluded_candidates", []), rows, "scope-review")):
            for number, candidate in enumerate(candidates):
                candidate_rule = rules.get(candidate['item_id'], {})
                prediction = project_reading_citations(candidate.get("judgment") or {})
                arithmetic_result = arithmetic.get((key, candidate['item_id']))
                arithmetic_withheld = False
                arithmetic_review_locations = []
                previous_arithmetic = (candidate.get('decision_trace') or {}).get('decision_source') == ARITHMETIC_METHOD or str(candidate.get('source') or '').endswith('#' + ARITHMETIC_METHOD)
                if (previous_arithmetic and (key, candidate['item_id']) in reading_payloads
                        and arithmetic_result is None and prediction.get('verdict') in {'COMPLIANT', 'VIOLATION'}):
                    arithmetic_withheld = True
                    arithmetic_review_locations = resolve_locations(document, prediction, evidence)
                    prediction = {**prediction, 'verdict': 'UNDETERMINED',
                                  'evidence_ids': [], 'evidence_line_refs': [], 'requirement_checks': [],
                                  'reason': '기존 자동 검산의 확정 판정을 보류했습니다. 원문의 금리 범위·적용 조건과 계산값의 대응이 현재 검산 요건을 충족하지 않습니다. 범위의 끝값을 계산 결과와 임의로 같다고 가정하거나, 범위 안에 있다는 이유만으로 전체 수치를 적정 처리할 수 없습니다.'}
                if arithmetic_result and not prediction.get('reading_quality_review'):
                    prediction = arithmetic_result
                applicability_audit = None
                applicability_review_locations = []
                reading_payload = reading_payloads.get((key, candidate['item_id']))
                if reading_payload and prediction:
                    guarded = copy.deepcopy(prediction)
                    guarded.setdefault('item_id', candidate['item_id'])
                    audit = apply_reading_guard(reading_payload, {'results': [guarded]})
                    if audit:
                        applicability_audit = {'verdict': prediction['verdict'],
                            'reason': prediction.get('reason', ''), 'status': 'WITHHELD_BY_READING_GUARD'}
                        prediction = guarded
                if prediction.get('verdict') in {'COMPLIANT', 'VIOLATION'}:
                    inconsistent = bool(reading_payload and observed_grounding_errors(
                        reading_payload, {**prediction, 'item_id': candidate['item_id']}))
                    for check in prediction.get('requirement_checks') or []:
                        if check.get('finding_basis') != 'OBSERVED' or check.get('status') not in {'SATISFIED', 'VIOLATED'}:
                            continue
                        refs = check.get('evidence_line_refs') or []
                        docs = [evidence[eid] for eid in check.get('evidence_ids', []) if eid in evidence]
                        # Legacy E-only citations can be checked against their
                        # saved source text, but never remapped to a guessed L.
                        window = cited_window_text(refs, evidence.values()) if refs else '\n'.join(
                            doc.get('text', '') for doc in docs)
                        if (refs or docs) and ungrounded_source_quotes(check.get('reason', ''), window):
                            inconsistent = True
                        if template_heading_only_citation(candidate_rule, check, list(evidence.values())):
                            inconsistent = True
                    if inconsistent:
                        applicability_audit = {'verdict': prediction['verdict'], 'reason': prediction.get('reason', ''),
                                               'status': 'WITHHELD_BY_GROUNDING_GUARD'}
                        prediction = {**prediction, 'verdict': 'UNDETERMINED',
                                      'evidence_ids': [], 'evidence_line_refs': [], 'requirement_checks': [],
                                      'reason': '인용한 원문이 판정 내용의 근거가 되지 않아 자동 확정을 보류했습니다. 제목만 인용했거나 실제 문구와 맞지 않는지 원본 본문을 확인해야 합니다.'}
                if unresolved_applicability(prediction):
                    applicability_review_locations = resolve_locations(document, prediction, evidence)
                    applicability_audit = {'verdict': prediction['verdict'], 'reason': prediction.get('reason', ''),
                                           'status': 'WITHHELD_BY_APPLICABILITY_GUARD'}
                    prediction = {**prediction, 'verdict': 'UNDETERMINED',
                                  'evidence_ids': [], 'evidence_line_refs': [], 'requirement_checks': [],
                                  'reason': '적용 여부를 확인하지 못한 항목이므로 미해당으로 제외하지 않고 사람 검토로 보냅니다. ' + prediction.get('reason', '')}
                if not prediction.get("verdict"):
                    continue  # Processing failures are reported separately, not model abstentions.
                item_id = candidate["item_id"]
                rule = {**rules.get(item_id, {}), **(rule_metadata or {}).get(item_id, {})}
                cited = [evidence[eid] for eid in prediction.get("evidence_ids", []) if eid in evidence]
                text = "\n".join(texts[ref] for ref in prediction.get("evidence_line_refs", []) if ref in texts)
                if not text:
                    text = "\n\n".join(doc.get("text", "") for doc in cited)
                title = rule.get("title") or item_id
                template_basis = candidate.get("template_basis") or {}
                if template_basis.get("requirement_mode") == "CONDITIONAL":
                    guidance = str(
                        ((template_basis.get("fields") or {}).get("guidance") or {}).get("text")
                        or ""
                    ).strip().splitlines()
                    if guidance:
                        title = f"{title} ({guidance[0]})"
                if ad.get("product_name"):
                    title = f"{ad['product_name']} · {title}"
                locations = resolve_locations(document, prediction, evidence)
                chunk_locations = resolve_chunk_locations(document, prediction, evidence)
                review_locations = (arithmetic_review_locations or applicability_review_locations
                                    or resolve_review_locations(document, prediction, evidence))
                location_status = "MAPPED" if locations else "NO_CITATION"
                if not locations and text:
                    location_status = "UNRESOLVED_REFERENCE"
                    if not any(page.get("canvas_w", 0) > 0 and page.get("canvas_h", 0) > 0
                               for page in document.get("pages", [])):
                        location_status = "SOURCE_GEOMETRY_MISSING"
                target.append({
                    "row_id": f"{key}:{kind}:{item_id}:{number}", "scope_id": key, "item_id": item_id,
                    "title": title, "question": rule.get("question", ""),
                    "criterion": rule.get("criterion") or rule.get("guide", ""),
                    "template_example": rule.get("example_text", "") if rule.get("source_sheet") == "HWPX_TEMPLATE" else "",
                    "template_appropriate_judgment": (
                        rule.get("appropriate_judgment", "")
                        if rule.get("source_sheet") == "HWPX_TEMPLATE" else ""
                    ),
                    "template_section": (candidate.get("template_basis") or {}).get("template_section"),
                    "template_requirement": (candidate.get("template_basis") or {}).get("requirement_mode"),
                    "requirement_checks": prediction.get("requirement_checks", []),
                    "verdict": {"VIOLATION": "위반", "COMPLIANT": "충족",
                                "UNDETERMINED": "판단불가", "NOT_APPLICABLE": "미해당"}.get(prediction["verdict"], prediction["verdict"]),
                    "reason": prediction.get("reason", ""), "evidence": text,
                    "evidence_ids": prediction.get("evidence_ids", []),
                    "evidence_line_refs": prediction.get("evidence_line_refs", []),
                    "evidence_locations": locations,
                    "chunk_locations": chunk_locations,
                    "review_locations": review_locations,
                    "evidence_location_status": location_status,
                    "reading_quality_review": prediction.get("reading_quality_review"),
                    "rule_basis": candidate.get("rule_basis"),
                    "decision_trace": ({'decision_source': ARITHMETIC_METHOD,
                        'model_result_source': candidate.get('source'),
                        'evidence_ids': prediction['evidence_ids'],
                        'evidence_line_refs': prediction['evidence_line_refs']}
                        if arithmetic_result and prediction is arithmetic_result else
                        {'decision_source': 'WITHHELD_BY_ARITHMETIC_GUARD',
                         'original_decision_source': ARITHMETIC_METHOD}
                        if arithmetic_withheld else candidate.get("decision_trace")),
                    "judgment_scope": "TEXT_ONLY" if (candidate.get("template_basis") or {}).get("text_facet_only") else "RULE",
                    "model_assessment": applicability_audit or ((reading_audits or {}).get((key, item_id))
                    if prediction.get("reading_quality_review") else None),
                })
    deferred = [dict(row, scope_id=ad.get("scope_id") or ad['ad_id']) for ad in ads
                for row in ad.get("deferred_rules", [])]
    scopes = {ad.get("scope_id") or ad["ad_id"] for ad in ads}
    for scope in (discovery or {}).get("ads", []):
        if scope.get("ad_id") in scopes:
            for source in scope.get('template_scope_deferred', []):
                if not any(row.get('scope_id') == scope['ad_id'] and row.get('item_id') == source.get('item_id')
                           and row.get('reason') == source.get('reason') for row in deferred):
                    deferred.append(dict(source, scope_id=scope['ad_id']))
            budget = scope.get("candidate_budget", {})
            # Discovery audit rows are outside the approved execution catalog;
            # they are diagnostics, not omitted judgments for the workspace.
            if (budget.get("method") != "discovery_audit_only_v1"
                    and budget.get("deferred_reason") != "outside_formal_execution_scope"):
                deferred.extend({"scope_id": scope["ad_id"], "item_id": item,
                                 "reason": "execution_budget_not_inapplicability"}
                                for item in budget.get("deferred_ids", []))
    for row in deferred:
        row["deferred_kind"] = (
            "EXECUTION_BUDGET" if row.get("reason") == "execution_budget_not_inapplicability"
            else "OTHER_TEMPLATE" if row.get("reason") in {"routing did not select this rule's template section",
                                                           "confirmed template does not match this v2 product subtype"}
            else "INPUT_OR_STRUCTURE"
        )
    discoveries = {ad['ad_id']: ad for ad in (discovery or {}).get('ads', [])}
    for entry in deferred:
        metadata = (rule_metadata or {}).get(entry['item_id'], {})
        for field in ('title', 'question', 'rule_basis'):
            if metadata.get(field):
                entry.setdefault(field, metadata[field])
        source = discoveries.get(entry['scope_id'], {})
        for field in ('template_candidates', 'presence_and_style', 'supplemental_v2_candidates', 'prohibition_candidates'):
            match = next((item for item in source.get(field, []) if item.get('item_id') == entry['item_id']), {})
            if match.get('title'):
                entry.setdefault('title', match['title'])
                break
    coverage = [ad.get('template_coverage') or discoveries.get(ad.get('scope_id') or ad['ad_id'], {}).get('template_coverage')
                for ad in ads]
    rows, excluded, omissions = consolidate_review_rows(
        rows, review_rows, deferred, scoped, coverage, document=document
    )
    reading_locations, reading_issues = local_reading_review(document)
    if reading_locations:
        rows.append({
            "row_id": f"{advertisement_id}:local-reading-review",
            "scope_id": advertisement_id,
            "item_id": "LOCAL_READING_REVIEW",
            "title": "원문 판독 확인",
            "question": "판독이 확정되지 않은 영역을 원본에서 확인했는가?",
            "criterion": "",
            "verdict": "판단불가",
            "reason": (f"광고 파일과 페이지 스캔은 완료됐지만 텍스트 판독이 확정되지 않은 "
                       f"위치 {len(reading_locations)}곳이 있습니다. 표시된 위치만 원본에서 확인해야 합니다."),
            "evidence": "",
            "evidence_ids": [],
            "evidence_line_refs": [],
            "evidence_locations": [],
            "chunk_locations": [],
            "review_locations": reading_locations,
            "evidence_location_status": "NO_CITATION",
            "reading_quality_review": {"policy": READING_QUALITY_VERSION, "issues": reading_issues},
            "rule_basis": None,
            "judgment_scope": "TEXT_ONLY",
            "model_assessment": None,
        })
    return {"rows": rows, "review_candidate_rows": [], "excluded_rows": excluded,
            "execution_omissions": omissions,
            "template_coverage": [value for value in coverage if value],
            "source_ads": ads,
            "output_failure_pairs": [pair for pair in raw.get("output_failure_pairs", [])
                                     if pair.get("ad_id") == advertisement_id],
            "deferred_rules": []}


def verified_separate_notice_lines(document, entry):
    """Resolve the narrow same-line prohibition from exact rendered line groups.

    This never estimates font, colour or visibility. It only accepts a complete
    exact-text mapping and verifies that separately marked notice paragraphs do
    not occupy the same rendered line. Ambiguous or unmarked structures remain
    manual.
    """
    basis = entry.get("template_basis") or {}
    fields = basis.get("fields") or {}
    guidance = str(
        basis.get("manual_guidance")
        or (fields.get("guidance") or {}).get("text")
        or entry.get("reason")
        or ""
    )
    if required_observation_medium(guidance)[2] != "원문줄구조":
        return None
    mappings = document.get("display_line_locations") or {}
    source_lines = {
        line["line_ref"]: str(line.get("text") or line.get("parser_text") or "").strip()
        for page in document.get("pages", [])
        for line in page_lines(page)
        if line.get("line_ref") and str(line.get("text") or line.get("parser_text") or "").strip()
    }
    if not source_lines or any(ref not in mappings for ref in source_lines):
        return None
    notices = {
        ref: text for ref, text in source_lines.items()
        if re.match(r"^\s*(?:※|\*|•|●)\s*\S", text)
    }
    if not notices:
        return None
    # Two independently marked notices inside one canonical line are already
    # ambiguous even before considering rendered coordinates.
    if any(len(re.findall(r"(?:^|\s)(?:※|\*|•|●)\s*\S", text)) > 1
           for text in notices.values()):
        return None
    occupied = []
    for ref in notices:
        for location in mappings[ref]:
            box = location.get("bbox") or []
            if len(box) != 4:
                return None
            occupied.append((ref, location.get("pageNo"), box))
    for index, (left_ref, left_page, left_box) in enumerate(occupied):
        for right_ref, right_page, right_box in occupied[index + 1:]:
            if left_ref == right_ref or left_page != right_page:
                continue
            overlap = min(left_box[3], right_box[3]) - max(left_box[1], right_box[1])
            if overlap > 0.5 * min(left_box[3] - left_box[1], right_box[3] - right_box[1]):
                return None
    return {"guidance": guidance, "notice_line_refs": list(notices),
            "notice_line_count": len(notices), "method": "EXACT_RENDERED_TEXT_LINES"}


def consolidate_review_rows(rows, review_rows, deferred, scoped, coverage, *, document=None):
    """One visible disposition per scope/item; source_ads retains raw history."""
    by_pair = {(row['scope_id'], row['item_id']): row for row in rows}
    for row in review_rows:
        row = dict(row, verdict='판단불가')
        by_pair.setdefault((row['scope_id'], row['item_id']), row)
    titles = {item['item_id']: item.get('title', '') for group in coverage if group
              for item in group.get('items', [])}
    omissions = []
    for entry in deferred:
        if entry['deferred_kind'] == 'OTHER_TEMPLATE':
            continue
        key, item = entry['scope_id'], entry['item_id']
        if entry['deferred_kind'] == 'EXECUTION_BUDGET':
            if not any(v['scope_id'] == key and v['item_id'] == item for v in omissions):
                omissions.append({'scope_id': key, 'item_id': item,
                                 'reason': '과거 실행 한도로 판정을 실행하지 못했습니다. 재처리가 필요합니다.'})
            reason = omissions[-1]['reason']
        else:
            reason = entry.get('reason') or '입력 또는 원문을 사람이 확인해야 합니다.'
            if reason == 'template source structure requires review':
                reason = '템플릿 원문 구조를 자동 해석할 수 없어 사람이 확인해야 합니다.'
        pair = (key, item)
        if pair not in by_pair:
            rule = scoped.get(key, ({}, {}))[0].get(item, {})
            basis = entry.get('template_basis') or {}
            by_pair[pair] = {'row_id': f'{key}:manual:{item}', 'scope_id': key, 'item_id': item,
                'title': entry.get('title') or titles.get(item) or rule.get('title') or '사람 확인이 필요한 점검항목',
                'question': entry.get('question') or rule.get('question', ''),
                'criterion': rule.get('criterion', ''), 'verdict': '판단불가', 'reason': '',
                'evidence': '', 'evidence_ids': [], 'evidence_line_refs': [],
                'evidence_locations': [], 'chunk_locations': [],
                'evidence_location_status': 'NO_CITATION', 'requirement_checks': [],
                'template_section': basis.get('template_section'),
                'template_requirement': basis.get('requirement_mode'),
                'rule_basis': entry.get('rule_basis'), 'judgment_scope': 'RULE'}
        row = by_pair[pair]
        # A visual rule has no model citation, but its retrieval trigger is the
        # exact place a reviewer should inspect. Keep that distinction in the
        # UI model: review_locations carries geometry while evidence_locations
        # remains reserved for automated judgment citations.
        trigger = next(iter(entry.get('trigger_evidence') or []), {}).get('trigger') or {}
        trigger_refs = list(dict.fromkeys(str(value) for value in trigger.get('line_refs') or []))
        if trigger_refs:
            trigger_locations = resolve_locations(
                document or {}, {'evidence_line_refs': trigger_refs}, {}
            )
            if trigger_locations:
                known = {value.get('key') for value in row.get('review_locations') or []}
                row['review_locations'] = [
                    *(row.get('review_locations') or []),
                    *(value for value in trigger_locations if value.get('key') not in known),
                ]
                row['evidence_line_refs'] = list(dict.fromkeys([
                    *(row.get('evidence_line_refs') or []), *trigger_refs,
                ]))
                if trigger.get('text') and not row.get('evidence'):
                    row['evidence'] = str(trigger['text'])
        line_observation = verified_separate_notice_lines(document or {}, entry)
        if line_observation and row.get("verdict") == "충족":
            row["judgment_scope"] = "RULE"
            row["line_structure_assessment"] = line_observation
            detail = "줄 단위 원본 확인: 각 유의사항이 서로 다른 줄에 있어 한 줄의 복수 유의사항 문구가 확인되지 않았습니다."
            if detail not in row.get("reason", ""):
                row["reason"] = (row.get("reason", "") + "\n" + detail).strip()
            continue
        reasons = row.setdefault('manual_review_reasons', [])
        if reason in reasons:
            continue
        reasons.append(reason)
        if row['verdict'] != '판단불가' and 'automated_assessment' not in row:
            row['automated_assessment'] = {'verdict': row['verdict'], 'reason': row['reason']}
        # A proven violation survives an unresolved visual facet. A text-only
        # pass is insufficient for an item whose other required facet is unknown.
        if row['verdict'] != '위반':
            row['verdict'] = '판단불가'
        row['reason'] = (row['reason'] + '\n사람 확인 필요: ' + reason).strip()
    return ([r for r in by_pair.values() if r['verdict'] != '미해당'],
            [r for r in by_pair.values() if r['verdict'] == '미해당'], omissions)


def frozen_rule_metadata(path, freeze):
    """Display names only from the exact regulation version frozen by this review."""
    path = Path(path)
    expected = tuple(sorted(item['sha256'] for item in freeze.get('inputs', [])
                            if str(item.get('path', '')).lower().endswith('.xlsx') and item.get('sha256')))
    if not expected or not path.is_file():
        return {}
    stat = path.stat()
    return _rule_metadata(str(path), stat.st_mtime_ns, stat.st_size, expected)


@lru_cache(maxsize=8)
def _rule_metadata(path, _mtime, _size, expected):
    import openpyxl
    digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if digest not in expected:
        return {}
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if '실행_점검항목' not in book.sheetnames:
            return {}
        rows = book['실행_점검항목'].iter_rows(values_only=True)
        columns = {name: index for index, name in enumerate(next(rows))}
        result = {}
        for row in rows:
            def field(name):
                return str(row[columns[name]] or '') if name in columns else ''
            item = field('항목ID')
            if item:
                basis = field('근거법령')
                result[item] = {'title': field('약칭'), 'question': field('점검문구'),
                    'rule_basis': {'item_id': item, 'source_type': 'REGULATION_V2',
                        'source_ref': '실행_점검항목:' + item, 'source_sha256': digest,
                        'legal_basis_refs': [basis] if basis else [],
                        'basis_status': 'LEGAL_BASIS_BOUND' if basis else 'LEGAL_BASIS_NOT_PROVIDED'}}
        return result
    finally:
        book.close()
