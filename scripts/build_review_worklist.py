"""Build a source-preserving worklist; design rows are never judgment inputs."""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def build(audit_path):
    review = json.loads(audit_path.read_text(encoding='utf-8'))
    canonical = json.loads((ROOT / 'rag-pipeline/config/canonical-execution-plans-v2.json').read_text(encoding='utf-8'))
    extra_statuses = {'미선정 우선 검토', '기존 보완에 잔여 결합', '미선정 사람 시각 검토'}
    rows = []
    for plan in canonical['plans']:
        if plan['source']['source_kind'] != 'TEMPLATE':
            continue
        source = plan['source']
        rows.append({'id': plan['plan_id'], 'label': source['label'], 'kind': 'TEMPLATE',
                     'scope': source.get('product_template'), 'status': 'CURRENT_TEMPLATE',
                     'selection': '필수 범위 전개 후 원문 조건·예외 확인',
                     'source': source, 'source_sha256': plan['source_sha256'],
                     'review_program': plan['review_program'], 'applicability_inputs': plan['applicability_inputs'],
                     'applicability_logic': plan['applicability_logic'], 'obligations': plan['obligations'],
                     'obligation_logic': plan['obligation_logic'],
                     'semantics_verified': False, 'execution_status': 'TEMPLATE_POLICY_CONNECTED',
                     'note': '출처·구조 연결을 검증한 현행 계획. 전 항목의 의미 정확도를 확정한 상태는 아님.'})
    selected = [r for r in review['reviews'] if r['current_selected'] or r['status'] in extra_statuses]
    assert len(selected) == 66 and sum(r['current_selected'] for r in selected) == 32
    for row in selected:
        status = row['status']
        if row['id'] in {'C-044', 'C-045'}:
            disposition = 'SCOPE_HOLD'
        elif row['id'] in {'C-137', 'D-194'}:
            disposition = 'SOURCE_VERSION_HOLD'
        elif row['id'] in {'C-085', 'C-081'}:
            disposition = 'POLICY'
        elif status in {'중복·복합항목 분해', '기존 보완에 잔여 결합'}:
            disposition = 'SHARED_WITH_RESIDUAL'
        elif status in {'사람 시각 검토', '미선정 사람 시각 검토'}:
            disposition = 'HUMAN_VISUAL'
        else:
            disposition = 'RESIDUAL_CHECK'
        source = row['original_v2']
        plan = next((p for p in canonical['plans'] if p['plan_id'] == row['id']), None)
        rows.append({'id': row['id'], 'label': row['label'], 'kind': 'SUPPLEMENT_32' if row['current_selected'] else 'ADDITIONAL_34',
                     'scope': source.get('적용상품', ''), 'status': disposition,
                     'selection': '선정 구조화 작업 대상; source scope·판본·예외와 잔여 검사 확정 후 별도 승격',
                     'source': source, 'source_refs': row['direct_regulatory_refs'],
                     'template_overlap_refs': row['current_template_refs'],
                     'owner': row['owner'], 'required_inputs': row['inputs'], 'risk': row['risk'], 'note': row['note'],
                     'existing_canonical_plan': plan, 'execution_status': 'DISABLED_DESIGN_WORKLIST',
                     'semantics_verified': False})
    counts = dict(Counter(r['status'] for r in rows if r['kind'] == 'ADDITIONAL_34'))
    remaining = [{key: r[key] for key in ('id', 'label', 'status', 'note')} for r in review['reviews'] if r not in selected]
    return {'schema_version': 'review-worklist-v1', 'date': '2026-09-28', 'operationally_connected': False,
            'purpose': '사이트에서 열람할 구조화 작업대장. 검색·판정 입력과 분리한다.',
            'source_audit_sha256': hashlib.sha256(audit_path.read_bytes()).hexdigest(),
            'canonical_source_binding_sha256': canonical['source_binding_sha256'],
            'support_scope': {'product_groups': ['예금성', '대출성', '투자성'],
                              'retirement_components': ['FUND', 'ETF', 'ELB'],
                              'unsupported': ['카드성', '보험', 'ETF·ELB 단독 광고', '퇴직연금 펀드 상세방법의 단독 광고 재사용', '공동광고·계열사 채무관계 자동 확인']},
            'counts': {'templates': 239, 'supplement_32': 32, 'additional_34': 34, 'work_rows': len(rows),
                       'additional_dispositions': counts, 'other_regulatory_rows': len(remaining)},
            'rows': rows, 'other_regulatory_rows': remaining,
            'promotion_requirements': ['원문 범위·판본·예외 확인', '템플릿과 공유할 사실·검사 및 잔여 검사 분리',
                                       '원자 검사·분기·허용결과 확정', '필요 입력·Rule/LLM/사람 담당 확정',
                                       '양성·음성·경계 검사', '정본 카탈로그·검색 게이트·화면·추적 연결 확인'],
            'gold_read': False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--audit', required=True, type=Path)
    parser.add_argument('--output', type=Path, default=ROOT / 'rag-pipeline/config/review-worklist-v1.json')
    parser.add_argument('--workbook', type=Path)
    args = parser.parse_args()
    value = build(args.audit)
    args.output.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    if args.workbook:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = '시작과 흐름'
        sheet.append(['단계', '작업', '주의'])
        for row in [('1', '상품·매체·운용상품 확인', '미확정은 조건 없음이 아니다'),
                    ('2', 'P1/P3 원문·라벨·줄·좌표 연결', '라벨·신뢰도 점수만으로 적정 확정 금지'),
                    ('3', '기본 템플릿 전수 전개 및 범위 확인', '필수 항목 선정에는 임베딩을 쓰지 않음'),
                    ('4', '광고→보완 후보 / 검사→광고 근거', '현재 보완 후보 검색·판정은 비활성'),
                    ('5', '원자 검사별 Rule + LLM + 사람·외부자료', '결론과 분기는 코드가 계산'),
                    ('6', '원문 인용·확인필요 및 독립 평가', '예측 동결 후 정답 대조')]:
            sheet.append(row)
        for kind, name in [('TEMPLATE', '템플릿239'), ('SUPPLEMENT_32', '기존보완32'), ('ADDITIONAL_34', '추가후보34')]:
            sheet = workbook.create_sheet(name)
            sheet.append(['ID', '항목', '범위', '상태', '담당', '필요 입력', '원문', '현재 검사 또는 기존 계획', '추가 검토', '판단기준 확정', '예외 확인', '양성', '음성', '경계'])
            for row in value['rows']:
                if row['kind'] != kind:
                    continue
                sheet.append([row['id'], row['label'], row['scope'], row['status'], row.get('owner', ''), row.get('required_inputs', ''),
                              json.dumps(row['source'], ensure_ascii=False),
                              json.dumps(row.get('obligations') or row.get('existing_canonical_plan') or {}, ensure_ascii=False), row['note'],
                              '', '', '', '', ''])
        sheet = workbook.create_sheet('기타규정190')
        sheet.append(['ID', '항목', '분류', '사유'])
        for row in value['other_regulatory_rows']:
            sheet.append([row['id'], row['label'], row['status'], row['note']])
        for sheet in workbook:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='145A41')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 28 if len(column) < 10 else 36
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for index in range(2, sheet.max_row + 1):
                sheet.row_dimensions[index].height = 80
        args.workbook.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(args.workbook)
    print(json.dumps(value['counts'], ensure_ascii=False))


if __name__ == '__main__':
    main()
